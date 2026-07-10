import io
from datetime import datetime
from zoneinfo import ZoneInfo

from flask import Response

TZ_PERU = ZoneInfo('America/Lima')

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                Paragraph, Spacer)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Color institucional (coincide con --rojo del CSS)
_ROJO = colors.HexColor('#8B0000')
_GRIS_BORDE = colors.HexColor('#e2e4e8')
_GRIS_FILA = colors.HexColor('#f4f5f7')


def _valor_celda(valor):
    if valor is None:
        return ''
    if isinstance(valor, datetime):
        return valor.strftime('%Y-%m-%d %H:%M')
    return str(valor)


def _nombre_archivo(nombre_base, extension):
    marca = datetime.now(TZ_PERU).strftime('%Y%m%d_%H%M')
    return "reporte_" + nombre_base + "_" + marca + "." + extension


def _valor_celda_xlsx(valor):
    # A diferencia de CSV/PDF, preservamos el tipo nativo cuando tiene sentido
    # (fechas y numeros) para que Excel los reconozca como tal (ordenar, sumar, etc.)
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.replace(tzinfo=None)
    if isinstance(valor, (int, float)):
        return valor
    return str(valor)


def generar_xlsx(nombre_base, columnas, filas):
    libro = Workbook()
    hoja = libro.active
    hoja.title = nombre_base[:31] or 'Reporte'

    relleno_header = PatternFill(start_color='8B0000', end_color='8B0000', fill_type='solid')
    fuente_header = Font(color='FFFFFF', bold=True)
    alineacion_header = Alignment(vertical='center')

    for i, col in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=i, value=col['titulo'])
        celda.fill = relleno_header
        celda.font = fuente_header
        celda.alignment = alineacion_header

    for fila_idx, fila in enumerate(filas, start=2):
        for col_idx, col in enumerate(columnas, start=1):
            valor = _valor_celda_xlsx(fila.get(col['clave']))
            celda = hoja.cell(row=fila_idx, column=col_idx, value=valor)
            if isinstance(valor, datetime):
                celda.number_format = 'DD/MM/YYYY HH:MM'

    # Ancho de columna aproximado segun el peso relativo (mismo criterio que el PDF)
    pesos = [col.get('peso', 1) for col in columnas]
    total_peso = sum(pesos) or 1
    for i, peso in enumerate(pesos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = max(10, round(peso / total_peso * 90))

    hoja.freeze_panes = 'A2'
    hoja.auto_filter.ref = hoja.dimensions

    buffer = io.BytesIO()
    libro.save(buffer)
    contenido = buffer.getvalue()
    buffer.close()

    return Response(
        contenido,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={
            'Content-Disposition': 'attachment; filename=' + _nombre_archivo(nombre_base, 'xlsx')
        }
    )


def generar_pdf(nombre_base, titulo, columnas, filas):
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1.2 * cm, rightMargin=1.2 * cm,
        topMargin=1.2 * cm, bottomMargin=1.2 * cm,
    )

    estilos = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        'TituloReporte', parent=estilos['Title'],
        fontSize=16, textColor=_ROJO, spaceAfter=4,
    )
    estilo_sub = ParagraphStyle(
        'SubReporte', parent=estilos['Normal'],
        fontSize=9, textColor=colors.HexColor('#6b7280'), spaceAfter=12,
    )
    estilo_celda = ParagraphStyle(
        'Celda', parent=estilos['Normal'], fontSize=7.5, leading=9,
    )
    estilo_celda_head = ParagraphStyle(
        'CeldaHead', parent=estilos['Normal'], fontSize=8,
        leading=10, textColor=colors.white, fontName='Helvetica-Bold',
    )

    elementos = []
    elementos.append(Paragraph(titulo, estilo_titulo))
    generado = datetime.now(TZ_PERU).strftime('%d/%m/%Y %H:%M')
    elementos.append(Paragraph(
        "ElSuper HelpDesk &nbsp;·&nbsp; Generado el " + generado
        + " &nbsp;·&nbsp; " + str(len(filas)) + " registro(s)",
        estilo_sub
    ))

    encabezado = [Paragraph(col['titulo'], estilo_celda_head) for col in columnas]
    datos = [encabezado]
    for fila in filas:
        datos.append([
            Paragraph(_valor_celda(fila.get(col['clave'])), estilo_celda)
            for col in columnas
        ])

    if len(filas) == 0:
        datos.append([Paragraph('Sin registros', estilo_celda)]
                     + ['' for _ in range(len(columnas) - 1)])

    # Ancho util = pagina - margenes. Repartimos segun el peso de cada columna
    # para que las celdas hagan wrap y nada se desborde del borde derecho.
    ancho_util = documento.width
    pesos = [col.get('peso', 1) for col in columnas]
    total_peso = sum(pesos)
    anchos = [ancho_util * peso / total_peso for peso in pesos]

    tabla = Table(datos, colWidths=anchos, repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), _ROJO),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, _GRIS_FILA]),
        ('GRID', (0, 0), (-1, -1), 0.4, _GRIS_BORDE),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
    ]))

    elementos.append(tabla)
    documento.build(elementos)

    pdf = buffer.getvalue()
    buffer.close()

    return Response(
        pdf,
        mimetype='application/pdf',
        headers={
            'Content-Disposition': 'attachment; filename=' + _nombre_archivo(nombre_base, 'pdf')
        }
    )


def generar_reporte(formato, nombre_base, titulo, columnas, filas):
    if formato == 'pdf':
        return generar_pdf(nombre_base, titulo, columnas, filas)
    return generar_xlsx(nombre_base, columnas, filas)
